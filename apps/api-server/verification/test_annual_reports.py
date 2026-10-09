"""Use the same strictly guarded disposable PostgreSQL database as report verification."""
import io
from pathlib import Path
from datetime import datetime, timedelta, timezone
from uuid import uuid4
import unittest
import asyncio

import test_report_results as fixture  # Validates *_test DB and migrates before importing app code.
from sqlalchemy import delete, func, select
from starlette.requests import Request
from openpyxl import Workbook
from app.db import engine, schema as db
from app.repositories.annual_reports import import_rows, parse_rows, backfill
from app.repositories.file_store import PostgresFileStore
from app.models.job import Job, JobStatus
from app.api.annual_reports import annual_reports, annual_history, export_annual_reports
from fastapi import HTTPException
from openpyxl import load_workbook
from urllib.parse import unquote

FILTERS = {'states': ['ASSAM'], 'rtos': ['UDALGURI - AS27'], 'fromYear': '2026', 'toYear': '2026',
           'categoryGroups': ['Two Wheeler'], 'fuels': ['PURE EV', 'ELECTRIC(BOV)']}
OBSERVED = datetime(2026, 10, 2, 9, 12, 13, tzinfo=timezone.utc)


def rows(headers, values):
    return [{'sheet': 'Sheet1', 'row_number': i, 'cells': cells} for i, cells in enumerate([headers, *values], 1)]


def request(owner='admin', role='admin'):
    result = Request({'type': 'http', 'headers': []})
    result.state.authenticated_role, result.state.authenticated_user = role, owner
    return result


class AnnualReportsTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        async with engine.begin() as connection:
            for table in (db.main_reports, db.report_update_history):
                await connection.execute(delete(table))
        self.files = []

    async def asyncTearDown(self):
        async with engine.begin() as connection:
            for table in (db.main_reports, db.report_update_history):
                await connection.execute(delete(table))
            await connection.execute(delete(db.stored_files).where(db.stored_files.c.id.in_(self.files)))
        await engine.dispose()

    async def ingest(self, source, data, owner='alpha', filters=None):
        async with engine.begin() as connection:
            return await import_rows(connection, source_key=source, name=source, rows=data,
                owner=owner, filters=filters or FILTERS, observed_at=OBSERVED)

    async def read(self, **kwargs):
        return await annual_reports(request(), year=kwargs.pop('year', 2026), dataset=kwargs.pop('dataset', ''),
            state=kwargs.pop('state', ''), rto=kwargs.pop('rto', ''), offset=kwargs.pop('offset', 0), limit=100)

    async def test_full_year_real_names_zero_missing_and_future_year(self):
        data = rows(['Maker', *[f'{m}’27' for m in ('JAN','FEB','MAR','APR','MAY','JUN','JUL','AUG','SEP','OCT','NOV','DEC')]],
                    [['ATHER ENERGY LTD', 0, 1, 2, 3, 4, 5, 6, 7, 8, None, 10, 11]])
        await self.ingest('year2027', data)
        result = await self.read(year=2027)
        self.assertEqual(result['years'], [2027])
        self.assertEqual(result['rows'][0]['months'], [0, 1, 2, 3, 4, 5, 6, 7, 8, None, 10, 11])
        self.assertEqual(result['rows'][0]['maker'], 'ATHER ENERGY LTD')
        self.assertEqual(result['rows'][0]['rto_code'], 'AS27')
        self.assertEqual(result['rows'][0]['rto'], 'UDALGURI')

    async def test_repeat_download_adds_only_missing_months_keeps_conflicting_zero(self):
        await self.ingest('first', rows(['Maker','2026-Jan','2026-Feb'], [['BAJAJ AUTO LTD',0,2]]))
        await self.ingest('second', rows(['Maker','2026-Jan','2026-Feb','2026-Mar'], [['BAJAJ AUTO LTD',9,2,3], ['TVS MOTOR COMPANY LTD',1,2,3]]))
        self.assertFalse(await self.ingest('second', rows(['Maker','2026-Jan'], [['BAJAJ AUTO LTD',99]])))
        result = await self.read()
        self.assertEqual(result['summary']['rows'], 2)
        self.assertEqual(result['rows'][0]['months'][:3], [0,2,3])
        history = await annual_history(request(), year=2026, dataset=result['datasetId'], state='', rto='', offset=0, limit=20)
        latest = history['rows'][0]
        self.assertEqual(history['total'], 2)
        self.assertEqual(latest['status'], 'review')
        self.assertEqual(latest['details']['newRows'], 1)
        self.assertEqual(latest['details']['newCells'], 4)
        self.assertEqual(latest['details']['conflicts'], 1)
        self.assertEqual(latest['observed_at'], OBSERVED)
        self.assertEqual(result['lastSaved']['newRows'],1)
        self.assertEqual(result['lastSaved']['newMonthValues'],4)
        self.assertEqual(result['lastSaved']['conflicts'],1)

    async def test_duplicate_workbooks_and_concurrent_imports_cannot_duplicate_rows(self):
        makers = [['TVS MOTOR COMPANY LTD',12], ['BAJAJ AUTO LTD',5]]
        data = rows(['Maker','2026-Jan'], makers)
        await asyncio.gather(self.ingest('parallel-a', data), self.ingest('parallel-b', rows(['Maker','2026-Jan'], makers[::-1])))
        result = await self.read()
        self.assertEqual(result['summary']['rows'], 2)
        async with engine.connect() as connection:
            self.assertEqual(await connection.scalar(select(func.count()).select_from(db.main_reports)), 2)
        await self.ingest('parallel-c', data)
        history = await annual_history(request(), year=2026, dataset=result['datasetId'], state='', rto='', offset=0, limit=20)
        self.assertEqual(history['rows'][0]['status'], 'unchanged')
        self.assertEqual(history['rows'][0]['details']['duplicates'], 2)

    async def test_other_names_invalid_counts_and_ambiguous_context_are_reported(self):
        data = rows(['Maker','2026-Jan','2026-Feb'], [
            ['Others',1,2], ['Other',3,4], ['Unknown',5,6],
            ['SUZUKI MOTORCYCLE INDIA PVT LTD','invalid',0]])
        await self.ingest('unresolved', data)
        result = await self.read()
        self.assertEqual(len(result['rows']), 2)
        others = next(row for row in result['rows'] if row['maker'] == 'OTHERS')
        suzuki = next(row for row in result['rows'] if row['maker'] == 'SUZUKI MOTORCYCLE INDIA PVT LTD')
        self.assertEqual(others['months'][:2], [1,2])
        self.assertEqual(suzuki['months'][:2], [None,0])
        history = await annual_history(request(), year=2026, dataset=result['datasetId'], state='', rto='', offset=0, limit=20)
        self.assertEqual(history['rows'][0]['details']['unresolvedMakers'], 2)
        self.assertEqual(history['rows'][0]['details']['invalidCells'], 1)
        entries, details = parse_rows(data, {**FILTERS, 'states': ['ASSAM', 'DELHI']})
        self.assertFalse(entries)
        self.assertEqual(details['missingContext'], 2)  # Both OTHERS and the named Maker need State/RTO.

    async def test_newer_crawl_replaces_values_and_preserves_previous_source(self):
        await self.ingest('old', rows(['Maker', '2026-Jan'], [['BAJAJ AUTO LTD', 7]]))
        async with engine.begin() as connection:
            await import_rows(connection, source_key='new', name='New crawl',
                rows=rows(['Maker', '2026-Jan'], [['BAJAJ AUTO LTD', 11]]), filters=FILTERS,
                owner='alpha', observed_at=OBSERVED + timedelta(seconds=2), update_newer=True, strict=True)
        result = await self.read()
        self.assertEqual(result['rows'][0]['months'][0], 11)
        self.assertEqual(result['lastSaved']['status'], 'updated')
        self.assertEqual(result['lastSaved']['updatedMonthValues'], 1)
        async with engine.connect() as connection:
            provenance = await connection.scalar(select(db.main_reports.c.month_sources))
        self.assertEqual(provenance['1']['replaced']['value'], 7)
        self.assertEqual(provenance['1']['replaced']['source']['sourceKey'], 'old')

    async def test_newer_identical_observation_blocks_late_changed_result(self):
        await self.ingest('initial', rows(['Maker', '2026-Jan'], [['BAJAJ AUTO LTD', 7]]))
        for source, seconds, value in [('newer-identical', 3, 7), ('late-older', 2, 11)]:
            async with engine.begin() as connection:
                await import_rows(connection, source_key=source, name=source,
                    rows=rows(['Maker', '2026-Jan'], [['BAJAJ AUTO LTD', value]]), filters=FILTERS,
                    owner='alpha', observed_at=OBSERVED + timedelta(seconds=seconds), update_newer=True, strict=True)
        self.assertEqual((await self.read())['rows'][0]['months'][0], 7)
        async with engine.connect() as connection:
            provenance = await connection.scalar(select(db.main_reports.c.month_sources))
        self.assertEqual(provenance['1']['sourceKey'], 'newer-identical')

    async def test_normal_worker_updates_main_table_while_another_case_is_running(self):
        from app.services import services
        job = await services.jobs.create(Job(runnerId='live-save-fixture', filters=FILTERS,
            status=JobStatus.WAITING_RESULT))
        waiting = await services.jobs.create(Job(runnerId='other-live-fixture', filters=FILTERS,
            status=JobStatus.WAITING_RESULT))
        await self.ingest('previous-crawl', rows(['Maker', '2026-Jan'], [['BAJAJ AUTO LTD', 7]]),
            filters=job.filters.model_dump(mode='json', by_alias=True))
        workbook = Workbook(); workbook.active.append(['Maker', '2026-Jan']); workbook.active.append(['BAJAJ AUTO LTD', 11])
        content = io.BytesIO(); workbook.save(content); workbook.close()
        saved = await PostgresFileStore().commit_excel(job.id, 'crawl.xlsx', content.getvalue(),
            observed_at=OBSERVED + timedelta(seconds=5), runner_id=job.runner_id)
        self.assertEqual(saved['status'], 'COMPLETED')
        self.assertEqual((await services.jobs.get(waiting.id)).status, JobStatus.WAITING_RESULT)
        self.assertEqual((await self.read())['rows'][0]['months'][0], 11)

    async def test_shared_scope_search_and_literal_wildcards(self):
        await self.ingest('alpha', rows(['Maker','2026-Jan'], [['BAJAJ AUTO LTD',5]]))
        await self.ingest('beta', rows(['Maker','2026-Jan'], [['BAJAJ AUTO LTD',8]]), owner='beta')
        one = await annual_reports(request('alpha','user'), year=2026, dataset='', state='ass', rto='AS27', offset=0, limit=100)
        self.assertEqual(one['rows'][0]['months'][0], 5)
        self.assertEqual(len(one['datasets']), 1)
        denied = await annual_reports(request('beta','user'), year=2026, dataset=one['datasetId'], state='', rto='', offset=0, limit=100)
        self.assertEqual(denied['rows'], one['rows'])
        self.assertFalse((await self.read(state='%'))['rows'])
        denied_history = await annual_history(request('beta','user'), year=2026, dataset=one['datasetId'], state='', rto='', offset=0, limit=20)
        self.assertEqual(denied_history['total'], 2)
        await self.ingest('other-fuel', rows(['Maker','2026-Jan'], [['BAJAJ AUTO LTD',100]]), filters={**FILTERS, 'fuels': ['PETROL']})
        self.assertEqual(len((await self.read())['datasets']), 2)

    async def test_all_sheets_upload_writes_only_main_table_and_repeated_data_is_unchanged(self):
        workbook = Workbook()
        for i, sheet in enumerate([workbook.active, workbook.create_sheet('Second')]):
            sheet.append(['S.No','STATE','RTO','RTOCode','Maker',"JAN'26","DEC'26"])
            sheet.append([i+1,'ASSAM','UDALGURI','AS27',['ATHER ENERGY LTD','BAJAJ AUTO LTD'][i],i,12])
        out=io.BytesIO(); workbook.save(out); workbook.close()
        store = PostgresFileStore()
        for _ in range(2):
            await store.put(name='manufacturers.xlsx', content=out.getvalue(), kind='upload',
                mime_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', extract=True)
        result = await self.read()
        self.assertEqual(result['summary']['rows'],2)
        self.assertEqual(result['rows'][0]['months'][0],0)
        self.assertEqual(result['rows'][0]['months'][11],12)
        history=await annual_history(request(),year=2026,dataset=result['datasetId'],state='',rto='',offset=0,limit=20)
        self.assertEqual(history['total'],2)
        self.assertEqual(history['rows'][0]['status'],'unchanged')
        async with engine.connect() as connection:
            self.assertEqual(await connection.scalar(select(func.count()).select_from(db.stored_files)),0)

    async def test_accounts_contribute_missing_values_and_concurrent_duplicates_to_one_table(self):
        await self.ingest('shared-alpha',rows(['Maker',"JAN'26"],[['BAJAJ AUTO LTD',0]]))
        await self.ingest('shared-beta',rows(['Maker',"JAN'26","FEB'26"],
            [['BAJAJ AUTO LTD',0,2],['TVS MOTOR COMPANY LTD',9,None]]),owner='beta')
        data=rows(['Maker',"MAR'26"],[['ATHER ENERGY LTD',3]])
        await asyncio.gather(self.ingest('parallel-alpha',data),self.ingest('parallel-beta',data,owner='beta'))
        views=[await annual_reports(request(owner,'user'),year=2026,dataset='',state='',rto='',offset=0,limit=100)
               for owner in ('alpha','beta','new-account')]
        self.assertEqual(views[0],views[1])
        self.assertEqual(views[1],views[2])
        self.assertEqual(views[0]['summary']['rows'],3)
        self.assertEqual(len(views[0]['datasets']),1)
        self.assertNotIn('owner_key',views[0]['datasets'][0])
        bajaj=next(row for row in views[0]['rows'] if row['maker']=='BAJAJ AUTO LTD')
        self.assertEqual(bajaj['months'][:2],[0,2])
        async with engine.connect() as connection:
            self.assertEqual(await connection.scalar(select(func.count()).select_from(db.main_reports)),3)
            owners=set(await connection.scalars(select(db.report_update_history.c.owner_key)))
            self.assertEqual(owners,{'alpha','beta'})

    async def test_shared_migration_merges_missing_months_preserves_zero_conflicts_and_dates(self):
        import copy
        import importlib.util
        from pathlib import Path
        from datetime import timedelta
        path=Path(__file__).resolve().parents[1]/'migrations/versions/0006_shared_main_reports.py'
        spec=importlib.util.spec_from_file_location('shared_migration',path)
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        base={'id':'first','owner_key':'alpha','scope_key':'old-alpha','scope_label':'EV','filters':FILTERS,
              'year':2026,'state':'ASSAM','rto':'UDALGURI','rto_code':'AS27','maker':'BAJAJ AUTO LTD',
              'created_at':OBSERVED,'updated_at':OBSERVED,'month_sources':{'1':{'sourceKey':'first','observedAt':OBSERVED.isoformat()}},
              **dict.fromkeys(module.MONTHS),'jan':0}
        other=copy.deepcopy(base)
        other.update(id='second',owner_key='beta',scope_key='old-beta',state='Assam',maker='Bajaj Auto Ltd',
            jan=9,feb=2,created_at=OBSERVED+timedelta(seconds=1),updated_at=OBSERVED+timedelta(seconds=2),
            month_sources={'1':{'sourceKey':'second','observedAt':OBSERVED.isoformat()},'2':{'sourceKey':'second'}})
        separate=copy.deepcopy(base);separate.update(id='petrol',filters={**FILTERS,'fuels':['PETROL']})
        future=copy.deepcopy(base);future.update(id='future',year=2027)
        original=copy.deepcopy([base,other,separate,future])
        merged,stats=module.consolidate(original)
        self.assertEqual(original,[base,other,separate,future])
        self.assertEqual(stats['afterRows'],3)
        self.assertEqual(stats['conflictingMonths'],1)
        common=next(row for row in merged if row['year']==2026 and row['filters']['fuels']==['ELECTRIC(BOV)','PURE EV'])
        self.assertEqual((common['jan'],common['feb'],common['mar']),(0,2,None))
        self.assertEqual(common['month_sources']['1']['mergedSources'][0]['value'],9)
        self.assertEqual(common['updated_at'],other['updated_at'])
        self.assertNotIn('owner_key',common)
        # Monthly commit order, rather than account or record creation order, selects the saved value.
        earlier,_=module.consolidate([base,other],{'second':OBSERVED-timedelta(seconds=1),'first':OBSERVED})
        self.assertEqual(earlier[0]['jan'],9)

    async def test_bare_months_require_single_year_and_totals_are_excluded(self):
        data = rows(['Maker','Jan','Feb','Total'], [['TVS MOTOR COMPANY LTD',1,2,3],['TOTAL',1,2,3]])
        entries, _ = parse_rows(data, FILTERS)
        self.assertEqual(len(entries), 1)
        entries, details = parse_rows(data, {**FILTERS,'toYear':'2027'})
        self.assertFalse(entries)
        self.assertEqual(details['issueCount'], 1)

    async def test_no_data_has_history_but_does_not_fabricate_maker_or_zeroes(self):
        async with engine.begin() as connection:
            await import_rows(connection, source_key='empty', name='No record found', rows=[], owner='alpha',
                filters=FILTERS, observed_at=OBSERVED, no_data=True)
        result = await self.read()
        self.assertFalse(result['rows'])
        history = await annual_history(request(), year=2026, dataset=result['datasetId'], state='', rto='', offset=0, limit=20)
        self.assertEqual(history['rows'][0]['status'], 'no-data')
        self.assertEqual(result['lastSaved']['status'],'no-data')
        self.assertEqual(result['lastSaved']['newRows'],0)

    async def test_last_saved_shows_repeated_filter_with_no_duplicate_rows_and_respects_search(self):
        data=rows(['Maker',"JAN'26","FEB'26"],[['TVS MOTOR COMPANY LTD',0,2]])
        await self.ingest('first-save',data)
        first=await self.read(rto='AS27')
        self.assertEqual(first['lastSaved']['status'],'added')
        await self.ingest('repeat-save',data,owner='another-account')
        latest=await self.read(rto='AS27')
        self.assertEqual(latest['summary']['rows'],1)
        self.assertEqual(latest['lastSaved']['status'],'unchanged')
        self.assertEqual(latest['lastSaved']['newRows'],0)
        self.assertEqual(latest['lastSaved']['alreadySavedMonthValues'],2)
        self.assertIsNone((await self.read(rto='AS99'))['lastSaved'])
        self.assertIsNone((await self.read(year=2027))['lastSaved'])

    async def test_export_filtered_reads_all_pages_real_name_null_zero_and_safe_text(self):
        values = [[f'MAKER {i:03d}', i, 0] for i in range(150)]
        values[0][0] = '=UNTRUSTED NAME'
        await self.ingest('export-alpha', rows(['Maker',"JAN'26","FEB'26"], values))
        await self.ingest('export-beta', rows(['Maker',"JAN'26"], [['CONTRIBUTED MAKER',99]]), owner='beta')
        await self.ingest('export-year', rows(['Maker',"JAN'27"], [['FUTURE MAKER',99]]))
        view = await annual_reports(request('alpha','user'),year=2026,dataset='',state='ass',rto='AS27',offset=0,limit=100)
        self.assertEqual(len(view['rows']),100)
        response = await export_annual_reports(request('alpha','user'),year=2026,dataset=view['datasetId'],
            state='ass',rto='AS27',confirm_all=False)
        self.assertEqual(response.headers['x-report-row-count'],'151')
        filename = unquote(response.headers['content-disposition'].split("UTF-8''",1)[1])
        self.assertEqual(filename,'Maker Month Wise Data of UDALGURI - AS27, ASSAM (2026).xlsx')
        workbook=load_workbook(io.BytesIO(Path(response.path).read_bytes()),data_only=False)
        sheet=workbook.active
        self.assertEqual(sheet.max_row,154)
        self.assertEqual(sheet.max_column,17)
        self.assertEqual([sheet.cell(3,c).value for c in range(6,18)], [m+"'26" for m in
            ('JAN','FEB','MAR','APR','MAY','JUN','JUL','AUG','SEP','OCT','NOV','DEC')])
        names=[sheet.cell(r,5).value for r in range(4,155)]
        self.assertEqual(names[:100],[row['maker'] for row in view['rows']])
        label_row=names.index('=UNTRUSTED NAME')+4
        self.assertEqual(sheet.cell(label_row,5).data_type,'s')
        self.assertEqual(sheet.cell(label_row,6).value,0)
        self.assertEqual(sheet.cell(label_row,7).value,0)
        self.assertIsNone(sheet.cell(label_row,17).value)
        self.assertIn('MAKER 149',names)
        self.assertEqual(sheet.freeze_panes,'F4')
        self.assertEqual(sheet.auto_filter.ref,'A3:Q154')
        workbook.close()
        Path(response.path).unlink(missing_ok=True)

    async def test_export_all_requires_confirmation_and_respects_shared_scope_and_year(self):
        await self.ingest('first',rows(['Maker',"JAN'26"],[['BAJAJ AUTO LTD',0]]))
        await self.ingest('other-office',rows(['Maker',"JAN'26"],[['TVS MOTOR COMPANY LTD',2]]),
            filters={**FILTERS,'rtos':['TINSUKIA - AS23']})
        await self.ingest('private',rows(['Maker',"JAN'26"],[['CONTRIBUTED MAKER',3]]),owner='beta')
        for state in ('','   '):
            with self.assertRaises(HTTPException) as error:
                await export_annual_reports(request('alpha','user'),year=2026,dataset='',state=state,rto='',confirm_all=False)
            self.assertEqual(error.exception.status_code,409)
        response=await export_annual_reports(request('alpha','user'),year=2026,dataset='',state='',rto='',confirm_all=True)
        self.assertEqual(response.headers['x-report-row-count'],'3')
        Path(response.path).unlink(missing_ok=True)
        self.assertIn('All%20RTOs%2C%20ASSAM%20%282026%29.xlsx',response.headers['content-disposition'])
        # Every account selects the same shared scope.
        alpha=(await annual_reports(request('alpha','user'),year=2026,dataset='',state='',rto='',offset=0,limit=100))['datasetId']
        for opts in ({'year':2027,'dataset':'','state':'','rto':''},
                     {'year':2026,'dataset':alpha,'state':'%','rto':''},
                     {'year':2026,'dataset':'missing-scope','state':'ASSAM','rto':''}):
            who=request('beta','user') if opts['state']=='ASSAM' else request('alpha','user')
            with self.assertRaises(HTTPException) as error:
                await export_annual_reports(who,**opts,confirm_all=True)
            self.assertEqual(error.exception.status_code,404)

    async def test_export_full_title_and_filename_match_port_blair_sample(self):
        values=[['BAJAJ AUTO LTD',2,0,1,0,9,0,24,1,0],
            ['SUZUKI MOTORCYCLE INDIA PVT LTD',0,1,1,1,0,0,0,2,0],
            ['TVS MOTOR COMPANY LTD',0,5,6,2,0,0,3,4,2]]
        await self.ingest('sample',rows(['Maker',*[m+"'26" for m in ('JAN','FEB','MAR','APR','MAY','JUN','JUL','AUG','SEP')]],values),
            filters={**FILTERS,'states':['Andaman & Nicobar Island'],'rtos':['Port Blair DTO - AN1']})
        response=await export_annual_reports(request('alpha','user'),year=2026,dataset='',state='',rto='AN1',confirm_all=False)
        filename=unquote(response.headers['content-disposition'].split("UTF-8''",1)[1])
        self.assertEqual(filename,'Maker Month Wise Data of Port Blair DTO - AN1, Andaman & Nicobar Island (2026).xlsx')
        workbook=load_workbook(io.BytesIO(Path(response.path).read_bytes()),data_only=True)
        sheet=workbook.active
        self.assertEqual(sheet['A1'].value,filename[:-5])
        self.assertEqual(sum(sheet.cell(r,c).value or 0 for r in range(4,7) for c in range(6,18)),64)
        self.assertEqual(sheet['F3'].fill.fgColor.rgb,'00FFC900')
        self.assertEqual(sheet['E3'].fill.fgColor.rgb,'0078AF4A')
        workbook.close()

    async def test_http_export_auth_confirmation_and_binary_excel_response(self):
        from unittest.mock import patch, AsyncMock
        await self.ingest('http',rows(['Maker',"JAN'26"],[['TVS MOTOR COMPANY LTD',3]]))
        path='/api/annual-reports/export?year=2026&state=ASSAM&rto=AS27'
        self.assertEqual((await fixture.asgi_request('GET',path))[0],401)
        with patch('app.main.authenticate_access_token',new=AsyncMock(return_value={
                'username':'alpha','role':'user','session_id':'fixture-session'})):
            headers={'authorization':'Bearer fixture-token'}
            code,_=await fixture.asgi_request('GET','/api/annual-reports/export?year=2026',headers=headers)
            self.assertEqual(code,409)
            code,content=await fixture.asgi_request('GET',path,headers=headers,binary=True)
            self.assertEqual(code,200)
            self.assertTrue(content.startswith(b'PK'))
            workbook=load_workbook(io.BytesIO(content),data_only=True)
            self.assertEqual(workbook.active['E4'].value,'TVS MOTOR COMPANY LTD')
            self.assertEqual(workbook.active['F4'].value,3)
            workbook.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
