from io import BytesIO
import asyncio
from dataclasses import replace
from uuid import uuid4
from zipfile import ZipFile

import pytest
from httpx import ASGITransport, AsyncClient

from app import excel_storage
from app.api import excel as excel_api
from app.config import settings
from app.main import application
from app.models.filters import VahanFilters
from app.models.job import Job, JobStatus
from app.realtime.runner_events import job_status
from app.services import services
from app.security import issue_access_token


RUNNER_HEADERS = {
    "X-VAHAN-RUNNER-ID": "test-runner",
    "X-VAHAN-RUNNER-TOKEN": settings.runner_token,
}
UI_HEADERS = {"Authorization": f"Bearer {issue_access_token(settings.ui_auth_username)}"}


@pytest.fixture(autouse=True)
def isolated_excel_reports(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_api, "_report_dir", tmp_path)
    monkeypatch.setattr(excel_storage, "settings", replace(excel_storage.settings, excel_report_dir=str(tmp_path)))


async def test_upload_and_download_excel_flow() -> None:
    # 1. Create a mock job with scenario name
    job = Job(
        id=uuid4(),
        runnerId="test-runner",
        scenarioName="Delhi EV / 2024: Two-Wheeler",
        status=JobStatus.WAITING_RESULT,
        filters=VahanFilters(yAxis="State Name"),
    )
    await services.jobs.create(job)

    # 2. Upload an Excel file
    excel_buffer = BytesIO()
    with ZipFile(excel_buffer, "w") as workbook:
        workbook.writestr("xl/workbook.xml", "<workbook/>")
    fake_excel_bytes = excel_buffer.getvalue()
    file_payload = {
        "file": ("temp.xlsx", BytesIO(fake_excel_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }

    async with AsyncClient(
        transport=ASGITransport(app=application),
        base_url="http://test",
    ) as client:
        upload_resp = await client.post(
            f"/api/jobs/{job.id}/upload-excel",
            files=file_payload,
            headers=RUNNER_HEADERS,
        )
        assert upload_resp.status_code == 200
        upload_data = upload_resp.json()
        assert upload_data["ok"] is True
        # Filename should be sanitized from scenarioName and have timestamp
        dt = job.created_at.astimezone() if job.created_at.tzinfo else job.created_at
        expected_name = f"Delhi EV _ 2024_ Two-Wheeler_{dt.strftime('%Y%m%d_%H%M%S')}_{str(job.id)[:8]}.xlsx"
        assert upload_data["fileName"] == expected_name
        assert upload_data["sizeBytes"] == len(fake_excel_bytes)

        # Update job to COMPLETED so it appears in reports list
        await services.jobs.update_status(job.id, JobStatus.COMPLETED)

        # Verify job updated in repo
        updated_job = await services.jobs.get(job.id)
        assert updated_job.excel_file_name == expected_name
        assert updated_job.excel_file_size == len(fake_excel_bytes)

        # 3. Download the Excel file
        from urllib.parse import unquote

        download_resp = await client.get(f"/api/jobs/{job.id}/excel", headers=UI_HEADERS)
        assert download_resp.status_code == 200
        assert download_resp.content == fake_excel_bytes
        assert expected_name in unquote(download_resp.headers.get("content-disposition", ""))

        # 4. List reports
        reports_resp = await client.get("/api/jobs/reports", headers=UI_HEADERS)
        assert reports_resp.status_code == 200
        reports = reports_resp.json()
        assert len(reports) >= 1
        found = next((r for r in reports if r["jobId"] == str(job.id)), None)
        assert found is not None
        assert found["scenarioName"] == "Delhi EV / 2024: Two-Wheeler"
        assert found["fileName"] == expected_name
        assert found["fileSize"] == len(fake_excel_bytes)

        verify_resp = await client.post("/api/jobs/reports/verify", json={"fileNames": [expected_name]}, headers=UI_HEADERS)
        assert verify_resp.status_code == 200
        assert verify_resp.json()["files"][expected_name] == len(fake_excel_bytes)

        stored_download = await client.get(f"/api/jobs/reports/file/{expected_name}", headers=UI_HEADERS)
        assert stored_download.status_code == 200
        assert stored_download.content == fake_excel_bytes


async def test_upload_rejects_non_excel_bytes() -> None:
    job = Job(
        id=uuid4(), runnerId="test-runner", status=JobStatus.WAITING_RESULT,
        filters=VahanFilters(yAxis="State Name"),
    )
    await services.jobs.create(job)
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        response = await client.post(
            f"/api/jobs/{job.id}/upload-excel",
            files={"file": ("report.xlsx", BytesIO(b"<html>error</html>"),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=RUNNER_HEADERS,
        )
    assert response.status_code == 400
    assert (await services.jobs.get(job.id)).excel_file_name is None


async def test_job_cannot_complete_without_saved_excel() -> None:
    job = Job(
        id=uuid4(), runnerId="test-runner", status=JobStatus.WAITING_RESULT,
        filters=VahanFilters(yAxis="State Name"),
    )
    await services.jobs.create(job)
    await services.runners.register(runner_id="test-runner", name="Test runner", socket_id="test-socket")
    await services.runners.set_job("test-runner", str(job.id))

    response = await job_status("test-socket", {"jobId": str(job.id), "status": "COMPLETED"})

    assert response["ok"] is False
    assert (await services.jobs.get(job.id)).status == JobStatus.WAITING_RESULT


async def test_no_data_keeps_state_and_rto_without_creating_excel(tmp_path) -> None:
    title = "Maker Month Wise Data  of Baratang - AN201 , Andaman & Nicobar Island (2026)"
    job = Job(
        id=uuid4(), runnerId="test-runner", status=JobStatus.WAITING_RESULT,
        scenarioName=title,
        filters=VahanFilters(states=["Andaman & Nicobar Island"], rtos=["Baratang - AN201"], yAxis="Maker"),
    )
    await services.jobs.create(job)
    await services.runners.register(runner_id="test-runner", name="Test runner", socket_id="test-socket")
    await services.runners.set_job("test-runner", str(job.id))

    response = await job_status("test-socket", {"jobId": str(job.id), "status": "NO_DATA"})

    stored = await services.jobs.get(job.id)
    runner = await services.runners.get("test-runner")
    assert response == {"ok": True}
    assert stored.status == JobStatus.NO_DATA
    assert stored.filters.states == ["Andaman & Nicobar Island"]
    assert stored.filters.rtos == ["Baratang - AN201"]
    assert stored.excel_file_name is None
    assert stored.no_data_file_name == f"{title}.txt"
    assert runner.current_job_id is None
    assert list(tmp_path.rglob("*.xlsx")) == []
    marker = next(tmp_path.rglob("*.txt"))
    assert marker.parent.name == "Andaman & Nicobar Island"
    assert marker.read_text(encoding="utf-8") == f"{title}\n"
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        downloaded = await client.get(f"/api/jobs/{job.id}/no-data", headers=UI_HEADERS)
    assert downloaded.status_code == 200
    assert downloaded.text == f"{title}\n"


async def test_cancelled_job_cannot_become_no_data_or_write_text(tmp_path) -> None:
    job = Job(
        id=uuid4(), runnerId="test-runner", status=JobStatus.WAITING_RESULT,
        scenarioName="Maker Month Wise Data  of Baratang - AN201 , Andaman & Nicobar Island (2026)",
        filters=VahanFilters(states=["Andaman & Nicobar Island"], rtos=["Baratang - AN201"]),
    )
    await services.jobs.create(job)
    await services.runners.register(runner_id="test-runner", name="Test runner", socket_id="test-socket")
    await services.runners.set_job("test-runner", str(job.id))

    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        cancelled = await client.post(f"/api/jobs/{job.id}/cancel", headers=UI_HEADERS)
    late_result = await job_status("test-socket", {"jobId": str(job.id), "status": "NO_DATA"})

    assert cancelled.status_code == 200
    assert (await services.jobs.get(job.id)).status == JobStatus.CANCELLED
    assert late_result["ok"] is False
    assert list(tmp_path.rglob("*.txt")) == []


async def test_upload_excel_non_existent_job() -> None:
    fake_job_id = uuid4()
    file_payload = {
        "file": ("report.xlsx", BytesIO(b"data"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    async with AsyncClient(
        transport=ASGITransport(app=application),
        base_url="http://test",
    ) as client:
        resp = await client.post(f"/api/jobs/{fake_job_id}/upload-excel", files=file_payload, headers=RUNNER_HEADERS)
        assert resp.status_code == 404


async def test_upload_excel_rejects_empty_file() -> None:
    job = Job(
        id=uuid4(),
        runnerId="test-runner",
        status=JobStatus.WAITING_RESULT,
        filters=VahanFilters(yAxis="State Name"),
    )
    await services.jobs.create(job)
    file_payload = {
        "file": ("empty.xlsx", BytesIO(b""), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }

    async with AsyncClient(
        transport=ASGITransport(app=application),
        base_url="http://test",
    ) as client:
        resp = await client.post(f"/api/jobs/{job.id}/upload-excel", files=file_payload, headers=RUNNER_HEADERS)

    assert resp.status_code == 400
    assert resp.json()["detail"] == "Excel file is empty."


async def test_download_excel_not_found() -> None:
    job = Job(
        id=uuid4(),
        runnerId="test-runner",
        status=JobStatus.COMPLETED,
        filters=VahanFilters(yAxis="State Name"),
    )
    await services.jobs.create(job)

    async with AsyncClient(
        transport=ASGITransport(app=application),
        base_url="http://test",
    ) as client:
        resp = await client.get(f"/api/jobs/{job.id}/excel", headers=UI_HEADERS)
        assert resp.status_code == 404


async def test_repeat_and_concurrent_uploads_save_one_readable_workbook(tmp_path) -> None:
    from openpyxl import Workbook, load_workbook

    job = await services.jobs.create(Job(
        runnerId="test-runner", status=JobStatus.WAITING_RESULT,
        scenarioName="Maker Month Wise Data  of Adoni RTO - AP221 , Andhra Pradesh (2026)",
        filters=VahanFilters(states=["Andhra Pradesh"], rtos=["Adoni RTO - AP221"], yAxis="Maker"),
    ))
    workbook = Workbook()
    workbook.active.append(["Maker", "2026-Jan", "Total"])
    workbook.active.append(["ATHER ENERGY LTD", 38, 487])
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    await services.runners.register(runner_id="test-runner", name="Test runner", socket_id="test-socket")
    await services.runners.set_job("test-runner", str(job.id))
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        async def upload(headers=RUNNER_HEADERS):
            return await client.post(
                f"/api/jobs/{job.id}/upload-excel", headers=headers,
                files={"file": ("report.xlsx", data, "application/octet-stream")},
            )

        first, second = await asyncio.gather(upload(), upload())
        assert first.status_code == second.status_code == 200
        assert first.json() == second.json()
        assert len(list(tmp_path.rglob("*.xlsx"))) == 1
        assert not list(tmp_path.rglob("*.tmp"))
        acknowledgement = await job_status("test-socket", {"jobId": str(job.id), "status": "COMPLETED"})
        assert acknowledgement == {"ok": True}
        assert (await services.runners.get("test-runner")).current_job_id is None
        retry = await upload()
        assert retry.status_code == 200 and retry.json() == first.json()
        wrong_runner = await upload({**RUNNER_HEADERS, "X-VAHAN-RUNNER-ID": "another-runner"})
        assert wrong_runner.status_code == 403
        downloaded = await client.get(f"/api/jobs/{job.id}/excel", headers=UI_HEADERS)
        assert downloaded.content == data
        readback = load_workbook(BytesIO(downloaded.content), read_only=True)
        assert list(readback.active.values)[1] == ("ATHER ENERGY LTD", 38, 487)
        readback.close()
        # A lost completion ACK can be retried after the next case is assigned.
        next_job = await services.jobs.create(Job(runnerId="test-runner", status=JobStatus.ASSIGNED, filters={}))
        await services.runners.set_job("test-runner", str(next_job.id))
        duplicate = await job_status("test-socket", {"jobId": str(job.id), "status": "COMPLETED"})
        assert duplicate == {"ok": True}
        assert (await services.runners.get("test-runner")).current_job_id == str(next_job.id)


async def test_storage_failure_returns_useful_error_with_cors(monkeypatch) -> None:
    job = await services.jobs.create(Job(
        runnerId="test-runner", status=JobStatus.WAITING_RESULT, filters={},
    ))

    def cannot_create(*_args, **_kwargs):
        raise PermissionError("disk permission denied")

    monkeypatch.setattr(excel_api, "report_session_dir", cannot_create)
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        response = await client.post(
            f"/api/jobs/{job.id}/upload-excel",
            headers={**RUNNER_HEADERS, "Origin": "http://127.0.0.1:5173"},
            files={"file": ("report.xlsx", b"test", "application/octet-stream")},
        )
    assert response.status_code == 500
    assert "storage and permissions" in response.json()["detail"]
    assert response.headers["access-control-allow-origin"] == "http://127.0.0.1:5173"
    assert (await services.jobs.get(job.id)).excel_file_name is None


async def test_cancellation_during_workbook_validation_leaves_no_report(tmp_path, monkeypatch) -> None:
    job = await services.jobs.create(Job(
        runnerId="test-runner", status=JobStatus.WAITING_RESULT, filters={},
    ))
    original_to_thread = asyncio.to_thread

    async def cancel_during_validation(function, *args):
        await original_to_thread(function, *args)
        await services.jobs.transition_status(job.id, JobStatus.CANCELLED)

    monkeypatch.setattr(excel_api.asyncio, "to_thread", cancel_during_validation)
    buffer = BytesIO()
    with ZipFile(buffer, "w") as workbook:
        workbook.writestr("xl/workbook.xml", "<workbook/>")
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        response = await client.post(
            f"/api/jobs/{job.id}/upload-excel", headers=RUNNER_HEADERS,
            files={"file": ("report.xlsx", buffer.getvalue(), "application/octet-stream")},
        )
    assert response.status_code == 409
    assert not list(tmp_path.rglob("*.xlsx"))
    assert not list(tmp_path.rglob("*.tmp"))
