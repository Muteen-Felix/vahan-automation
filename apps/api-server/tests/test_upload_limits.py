"""Small hostile archives must fail before dimensions expand into huge rows."""
from io import BytesIO
from zipfile import ZipFile

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import func, select

from app.db import engine, schema as db
from app.main import application
from app.models.job import Job, JobStatus
from app.repositories.file_store import extract_file
from app.services import services
from app.config import settings


def inflated_workbook(dimension):
    workbook = Workbook()
    workbook.active.append(['Maker', 'JAN', 'FEB', 'TOTAL'])
    workbook.active.append(['Fixture Maker', 5, 7, 12])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    rewritten = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(rewritten, 'w') as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                content = content.replace(b'ref="A1:D2"', f'ref="{dimension}"'.encode())
            target.writestr(entry, content)
    return rewritten.getvalue()


@pytest.mark.parametrize('dimension', ['A1:XFD1048576', 'A1:D1048576', 'A1:SR10000'])
async def test_inflated_sheet_dimensions_are_rejected_without_saving(dimension, admin_token):
    job = await services.jobs.create(Job(runnerId='dimension-runner', ownerUsername='test-admin',
        status=JobStatus.WAITING_RESULT, filters={'states':['Fixture State'], 'rtos':['Office - TS1'],
        'fromYear':'2025', 'toYear':'2025', 'yAxis':'Maker'}))
    async with AsyncClient(transport=ASGITransport(app=application), base_url='http://test') as client:
        response = await client.post(f'/api/jobs/{job.id}/main-report',
            headers={'X-VAHAN-RUNNER-ID':'dimension-runner', 'X-VAHAN-RUNNER-TOKEN':settings.runner_token},
            files={'file':('report.xlsx', inflated_workbook(dimension))})
    assert response.status_code == 400
    assert 'exceeds' in response.json()['detail']
    assert (await services.jobs.get(job.id)).status == JobStatus.WAITING_RESULT
    async with engine.connect() as connection:
        assert await connection.scalar(select(func.count()).select_from(db.main_reports)) == 0
        assert await connection.scalar(select(func.count()).select_from(db.report_update_history)) == 0


async def test_general_upload_returns_client_error_for_broken_zip(admin_headers):
    async with AsyncClient(transport=ASGITransport(app=application), base_url='http://test', headers=admin_headers) as client:
        response = await client.post('/api/files', files={'file':('broken.xlsx', b'not-a-zip')})
    assert response.status_code == 400


async def test_general_upload_returns_client_error_for_oversized_csv_field(admin_headers):
    async with AsyncClient(transport=ASGITransport(app=application),base_url='http://test',headers=admin_headers) as client:
        response = await client.post('/api/files',files={'file':('broken.csv',b'x'*200_000)})
    assert response.status_code == 400


async def test_sheet_cell_budget_is_shared_across_worksheets(monkeypatch):
    from app.repositories import file_store
    workbook = Workbook()
    for sheet in [workbook.active, workbook.create_sheet()]:
        sheet.append(['A', 'B', 'C'])
        sheet.append([1, 2, 3])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    monkeypatch.setattr(file_store, 'MAX_EXTRACTED_CELLS', 10)
    with pytest.raises(ValueError, match='cells'):
        extract_file(output.getvalue(), 'report.xlsx')


async def test_csv_width_is_bounded():
    with pytest.raises(ValueError, match='cell limit'):
        extract_file((','.join(['cell'] * 513)).encode(), 'report.csv')


async def test_api_data_and_authentication_failures_are_not_cached(admin_headers):
    async with AsyncClient(transport=ASGITransport(app=application),base_url='http://test') as client:
        denied = await client.get('/api/user-state')
        allowed = await client.get('/api/user-state',headers=admin_headers)
    assert denied.status_code == 401 and allowed.status_code == 200
    for response in [denied,allowed]:
        assert response.headers['Cache-Control'] == 'no-store'
        assert response.headers['X-Content-Type-Options'] == 'nosniff'
