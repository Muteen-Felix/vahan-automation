from __future__ import annotations
import asyncio, base64, csv, hashlib, io, json, zipfile
from datetime import date, datetime
from uuid import uuid4
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from openpyxl import load_workbook
from app.db import engine
from app.db import schema as db
from app.models.job import Job, JobStatus
from app.repositories.postgres import now, save_job, release_runner

MAX_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_EXTRACTED_ROWS = 500_000

class ExtractionLimit(ValueError):
    pass

def cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value if isinstance(value, (str, int, float, bool)) or value is None else str(value)

def extract_file(content, name):
    """Read every worksheet in memory before committing the main report."""
    rows, info = [], {}
    if name.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("The file is not an Excel workbook.")
            if len(archive.infolist()) > 10_000 or sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ExtractionLimit("Expanded workbook exceeds the parser limit.")
            if archive.testzip() is not None:
                raise ValueError("The workbook is damaged.")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            info["sheets"] = []
            for sheet in workbook:
                count = 0
                for index, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                    rows.append({"sheet": sheet.title, "row_number": index, "cells": [cell_value(v) for v in cells]})
                    count += 1
                    if len(rows) > MAX_EXTRACTED_ROWS:
                        raise ExtractionLimit("Workbook exceeds 500,000 extracted rows.")
                info["sheets"].append({"name": sheet.title, "rows": count})
        finally:
            workbook.close()
    elif name.lower().endswith(".csv"):
        for index, values in enumerate(csv.reader(io.StringIO(content.decode("utf-8-sig"))), 1):
            rows.append({"sheet": "CSV", "row_number": index, "cells": values})
            if len(rows) > MAX_EXTRACTED_ROWS:
                raise ExtractionLimit("CSV exceeds the parser row limit.")
    elif name.lower().endswith(".json"):
        info["document"] = json.loads(content)
    elif name.lower().endswith((".txt", ".log")):
        info["text"] = content.decode("utf-8")
    info["rowCount"] = len(rows)
    return rows, info

async def insert_file(connection, *, name, content, kind, mime_type, job=None, owner=None, rows=None, metadata=None):
    if kind in {'excel', 'no-data'}:
        raise ValueError('Report copies have been retired; use the main report transaction.')
    file_id = str(uuid4())
    values = dict(id=file_id, name=name, content=content, kind=kind, mime_type=mime_type,
        job_id=str(job.id) if job else None, owner_username=job.owner_username if job else owner,
        size=len(content), sha256=hashlib.sha256(content).hexdigest(), metadata=metadata or {}, created_at=now())
    inserted = await connection.scalar(pg_insert(db.stored_files).values(**values).on_conflict_do_nothing(
        constraint="uq_job_file").returning(db.stored_files.c.id))
    if not inserted:
        return await connection.scalar(select(db.stored_files.c.id).where(db.stored_files.c.job_id == str(job.id),
            db.stored_files.c.kind == kind, db.stored_files.c.name == name))
    return file_id

class PostgresFileStore:
    async def put(self, *, name, content, kind, mime_type, job=None, owner=None, metadata=None, extract=False):
        if kind in {'excel', 'upload'} and extract:
            rows, _ = await asyncio.to_thread(extract_file, content, name)
            source = 'upload:' + str(uuid4())
            from app.repositories.annual_reports import import_rows
            async with engine.begin() as connection:
                await import_rows(connection, source_key=source, name=name, rows=rows,
                    filters=job.filters.model_dump(mode='json', by_alias=True) if job else {},
                    owner=job.owner_username if job else owner, observed_at=now(), strict=True,
                    checksum=hashlib.sha256(content).hexdigest())
                summary = await connection.scalar(select(db.report_update_history.c.details)
                    .where(db.report_update_history.c.source_key == source))
            return {'sourceKey': source, 'name': name, 'summary': summary}
        if kind in {'excel', 'no-data'}:
            raise ValueError('Reports must be committed directly to the main table.')
        async with engine.begin() as connection:
            file_id = await insert_file(connection, name=name, content=content, kind=kind,
                mime_type=mime_type, job=job, owner=owner, metadata=metadata)
        return await self.get(file_id, include_content=False)

    async def commit_excel(self, job_id, name, content, *, observed_at=None, page_url='', runner_id=None):
        rows, _ = await asyncio.to_thread(extract_file, content, name)
        checksum = hashlib.sha256(content).hexdigest()
        async with engine.begin() as connection:
            payload = await connection.scalar(select(db.jobs.c.payload).where(db.jobs.c.id == str(job_id)).with_for_update())
            if not payload:
                raise ValueError('Job not found.')
            job = Job.model_validate(payload)
            if runner_id is not None and runner_id != job.runner_id:
                raise PermissionError('This runner is not assigned to the job.')
            if job.main_report_checksum:
                if job.main_report_checksum != checksum or job.status != JobStatus.COMPLETED:
                    raise ValueError('A different report has already been saved for this job.')
                return {'ok': True, 'status': job.status.value, 'jobId': str(job.id), 'summary': job.main_report_summary}
            if job.status != JobStatus.WAITING_RESULT:
                raise ValueError('Job is not waiting for a main report.')
            source_key = 'job:' + str(job.id)
            observed_at = observed_at or now()
            from app.repositories.annual_reports import import_rows
            await import_rows(connection, source_key=source_key, job_id=str(job.id), name=name, rows=rows,
                filters=job.filters.model_dump(mode='json', by_alias=True), owner=job.owner_username,
                observed_at=observed_at, strict=True, checksum=checksum)
            summary = await connection.scalar(select(db.report_update_history.c.details)
                .where(db.report_update_history.c.source_key == source_key))
            job.status = JobStatus.COMPLETED
            job.main_report_saved_at, job.main_report_checksum, job.main_report_summary = now(), checksum, summary
            job.result_message, job.result_observed_at = 'Data saved to main table', observed_at
            job.report_row_count, job.report_table_count = summary['parsedRows'], 1
            job.excel_file_name = job.excel_file_size = job.no_data_file_name = None
            job.error = None
            job.touch()
            await save_job(connection, job, 'main-report-saved')
            await release_runner(connection, job.runner_id, job.id)
        return {'ok': True, 'status': job.status.value, 'jobId': str(job.id), 'summary': summary}

    async def get(self, file_id, include_content=True):
        columns = list(db.stored_files.c) if include_content else [c for c in db.stored_files.c if c.name != "content"]
        async with engine.connect() as connection:
            row = (await connection.execute(select(*columns).where(db.stored_files.c.id == str(file_id)))).mappings().first()
        return dict(row) if row else None

    async def for_job(self, job_id, kind):
        async with engine.connect() as connection:
            row = (await connection.execute(select(db.stored_files).where(db.stored_files.c.job_id == str(job_id),
                db.stored_files.c.kind == kind).order_by(db.stored_files.c.created_at.desc()).limit(1))).mappings().first()
        return dict(row) if row else None

    async def list(self, owner=None, kind=None):
        query = select(*[c for c in db.stored_files.c if c.name != "content"]).order_by(db.stored_files.c.created_at.desc())
        if owner:
            query = query.where(db.stored_files.c.owner_username == owner)
        if kind:
            query = query.where(db.stored_files.c.kind == kind)
        async with engine.connect() as connection:
            return [dict(row) for row in (await connection.execute(query)).mappings()]

    async def rows(self, file_id, offset=0, limit=100):
        raise ValueError('Extracted file copies have been retired. Read the main report table.')

class PostgresCaptchaStore:
    async def save(self, job_id, data_url):
        from app.repositories.postgres import PostgresJobRepository
        header, encoded = data_url.split(",", 1)
        if header not in {"data:image/png;base64", "data:image/jpeg;base64", "data:image/gif;base64", "data:image/webp;base64"}:
            raise ValueError("Unsupported CAPTCHA image.")
        try:
            content = base64.b64decode(encoded, validate=True)
        except Exception as error:
            raise ValueError("Invalid CAPTCHA bytes.") from error
        if not content or len(content) > 750_000:
            raise ValueError("Invalid CAPTCHA image size.")
        job = await PostgresJobRepository().get(job_id)
        if not job:
            raise ValueError("Unknown job.")
        extension = header.split("/")[1].split(";")[0]
        result = await PostgresFileStore().put(name=f"captcha-{hashlib.sha256(content).hexdigest()}.{extension}",
            content=content, kind="captcha", mime_type=header[5:].split(";")[0], job=job)
        return result["id"]
