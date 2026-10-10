from pathlib import Path
import ast, json, re
import shutil
from html.parser import HTMLParser
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path('/Users/mac/Desktop/vahan-automation')
WORK=ROOT/'tmp/pdfs/baocaoNew2-update'
doc=Document()
sec=doc.sections[0]
sec.page_width=Pt(595.276); sec.page_height=Pt(841.89)
sec.top_margin=Pt(52);sec.bottom_margin=Pt(52)
sec.left_margin=Pt(44);sec.right_margin=Pt(44)
sec.header_distance=Pt(18);sec.footer_distance=Pt(18)
for name,size in [('Normal',9.4),('Title',21),('Heading 1',13),('Heading 2',13),('Caption',8.2)]:
 s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.line_spacing=Pt(13.3 if name=='Normal' else size+4)
 if name.startswith('Heading'):
  s.font.bold=True;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(9)
  s.paragraph_format.keep_with_next=True
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Caption'].font.bold=False
for style in doc.styles:
 for border in style.element.xpath('./w:pPr/w:pBdr'):
  border.getparent().remove(border)

class RichText(HTMLParser):
 def __init__(self,paragraph): super().__init__();self.paragraph=paragraph;self.bold=0
 def handle_starttag(self,tag,attrs):
  if tag=='b': self.bold+=1
  if tag=='br': self.paragraph.add_run().add_break()
 def handle_endtag(self,tag):
  if tag=='b': self.bold=max(0,self.bold-1)
 def handle_data(self,data):
  r=self.paragraph.add_run(data)
  if self.bold:r.bold=True

def para(text,style='Normal'):
 p=doc.add_paragraph(style=style);RichText(p).feed(text)
 return p
def heading(text): return para(text,'Heading 2' if text.startswith('2.1.') else 'Heading 1')
def page(): doc.add_page_break()
diagrams=json.loads((WORK/'diagrams.json').read_text())
inserted=set()
def figure(number):
 d=diagrams[number-1]
 caption=para(d['title'],'Caption')
 caption.paragraph_format.keep_with_next=True
 caption.paragraph_format.space_before=Pt(5)
 for r in caption.runs:r.bold=True
 path=WORK/'diagrams'/f"{d['id']}.png"
 with Image.open(path) as im:iw,ih=im.size
 scale=min(507.276/iw,(360 if number==1 else 480)/ih)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1
 p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(path),width=Pt(iw*scale),height=Pt(ih*scale))
 for inline in p._p.xpath('.//wp:docPr'):
  inline.set('descr',d['title']+'. '+d['note'])
 note=para(d['note'],'Caption')
 note.paragraph_format.space_after=Pt(10)
 inserted.add(number)
def shade(cell,fill):
 el=OxmlElement('w:shd');el.set(qn('w:fill'),fill);cell._tc.get_or_add_tcPr().append(el)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.LEFT;t.autofit=False
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders)
 for col,width in zip(t.columns,widths): col.width=Pt(width)
 for i,row_data in enumerate([headers]+rows):
  row=t.rows[0] if i==0 else t.add_row()
  rp=row._tr.get_or_add_trPr()
  rp.append(OxmlElement('w:cantSplit'))
  if i==0: rp.append(OxmlElement('w:tblHeader'))
  for j,value in enumerate(row_data):
   cell=row.cells[j];cell.width=Pt(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for side,twips in [('top',100),('bottom',100),('left',140),('right',140)]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),str(twips));e.set(qn('w:type'),'dxa');m.append(e)
   pr.append(m)
   p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
   p.paragraph_format.line_spacing=Pt(11.5)
   r=p.add_run(value);r.font.size=Pt(8.5);r.font.name='Arial'
   if i==0:
    shade(cell,'1D5276');r.font.color.rgb=RGBColor(255,255,255);r.bold=True
   elif i%2==0:shade(cell,'F4F7F9')
 return t

h=sec.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=h.add_run('VAHAN AUTOMATION | CUSTOMER REPORT');r.font.size=Pt(8)
f=sec.footer.paragraphs[0];f.paragraph_format.space_after=Pt(0)
f.paragraph_format.tab_stops.clear_all()
f.paragraph_format.tab_stops.add_tab_stop(Pt(507.276),WD_TAB_ALIGNMENT.RIGHT)
r=f.add_run('Current-source revision | 10 October 2026\tPage ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

tree=ast.parse((WORK/'build_report.py').read_text())
env={'W':507.276}
def value(n): return eval(compile(ast.Expression(n),'report-content','eval'),{'__builtins__':{}},env)
active=False
pending=None
for node in tree.body:
 if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id=='diagrams' for t in node.targets):break
 if not isinstance(node,ast.Expr) or not isinstance(node.value,ast.Call):continue
 call=node.value
 if isinstance(call.func,ast.Attribute) and isinstance(call.func.value,ast.Name) and call.func.value.id=='story' and call.func.attr=='append':
  inner=call.args[0]
  if isinstance(inner,ast.Call) and isinstance(inner.func,ast.Name) and inner.func.id=='p':
   args=[value(x) for x in inner.args]
   if args[1]=='TitleReport': active=True;para(args[0],'Title')
 elif active and isinstance(call.func,ast.Name):
  name=call.func.id
  if name=='para':
   text=value(call.args[0])
   if text.startswith('All diagrams are placed after the original report sections.'):continue
   if text.startswith('This report explains how VAHAN reports are collected, validated, saved and shared.'):
    text='This report explains how VAHAN reports are collected, validated, saved and shared. It retains the original section order (1-10 and 2.1) and describes the current system components. English diagrams appear in their corresponding sections alongside the explanation and tables.'
   text=text.replace('Figure 1 in Appendix A shows this architecture.','Figure 1 shows this architecture.')
   text=text.replace('Figure 6 shows the lifecycle boundary.','Figure 6 shows the lifecycle boundary.')
   paragraph=para(text)
   match=re.search(r'Figure ([1-9])\b',text)
   if match and int(match.group(1)) not in inserted:
    figure(int(match.group(1)))
  elif name=='heading':
   text=value(call.args[0])
   if text.startswith('Appendix A.'):continue
   if pending:
    figure(pending);pending=None
   heading(text)
  elif name=='page':pass
  elif name=='table':table(*[value(x) for x in call.args])

if pending:figure(pending)
assert inserted==set(range(1,10)),inserted

doc.core_properties.title='VAHAN System Overview Report'
doc.core_properties.subject='Current system report with editable text and tables'
doc.core_properties.author='VAHAN Automation'
doc.core_properties.keywords='VAHAN, system report, profiles, schedules, PostgreSQL, browser workers'
target=ROOT/'outputs/baocaoNew2_updated.docx'
backup=WORK/'baocaoNew2_before_inline.docx'
if target.exists() and not backup.exists():shutil.copy2(target,backup)
doc.save(target)
check=Document(ROOT/'outputs/baocaoNew2_updated.docx')
assert len(check.tables)==12,len(check.tables)
assert len(check.inline_shapes)==9
assert not check.settings.element.xpath('.//w:documentProtection')
print('Saved editable Word report:',len(check.paragraphs),'paragraphs,',len(check.tables),'native tables,',len(check.inline_shapes),'diagram images')
